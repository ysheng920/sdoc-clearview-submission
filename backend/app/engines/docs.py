"""Turn an attachment into plain text, plus labelled lines with char offsets.

Offsets matter: every extracted value carries a Span back into this exact text,
and the UI highlights it. Any reader that loses offsets breaks the highlighting,
so each reader returns the same text the spans index into.
"""
from __future__ import annotations

import logging
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

# pypdf narrates every malformed xref on stderr; the recovery is what matters.
logging.getLogger("pypdf").setLevel(logging.ERROR)

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# "Label: value" -- value must be non-empty, label short enough to be a label.
#
# The label accepts any character except the colon that ends it. An explicit
# whitelist looked safer and silently lost 8 documents: pypdf substitutes
# U+25A0 for glyphs it cannot map, so "TOTAL Gross Weight■■(KGS): 143,940 KG"
# failed to match and the only total weight on the page went unread. Nothing
# is put at risk by being permissive here -- field_for_label decides what is
# actually a field, and it normalises the junk away.
_LABEL_LINE = re.compile(r"^[ \t]*([^\W\d_][^\r\n:：]{1,48}?)[ \t]*[:：][ \t]*(\S.*?)[ \t]*$", re.M)

# A line that is only a label: nothing follows it, and it is short enough to be one.
_LABEL_ONLY = re.compile(r"^[ \t]*([^\W\d_][^\r\n:：]{1,58}?)[ \t]*[:：]?[ \t]*$")


# A photographed or screenshotted document. It has no text layer by definition,
# so it goes straight to the same vision path a scanned PDF takes. Reading one
# as text used to return replacement characters rather than raise, which meant
# label parsing ran on binary noise and the page was never actually looked at.
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


class UnreadableDocument(Exception):
    """Raised when no text could be recovered -- the caller escalates to a human."""


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        raise UnreadableDocument(f"{path.name}: an image has no text layer")
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    # python-docx would work but the format is a zip of XML; stdlib is enough.
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    lines = []
    for para in root.iter(f"{_W_NS}p"):
        text = "".join(node.text or "" for node in para.iter(f"{_W_NS}t"))
        lines.append(text)
    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue
            # Two-column sheets are label/value pairs; rebuild them as such so the
            # same label regex works on spreadsheets and text files alike.
            lines.append(f"{cells[0]}: {cells[1]}" if len(cells) == 2 else "  ".join(cells))
    wb.close()
    return "\n".join(lines)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if not text.strip():
        raise UnreadableDocument(f"{path.name}: no extractable text (likely a scan)")
    return text


def labelled_lines(text: str, is_label=None) -> list[tuple[str, str, int, int]]:
    """Every label/value pair as (label, value, value_start, value_end).

    Two layouts appear in this corpus and only the first was handled at first:

        Shipper/Exporter: APRIL FAR EAST (M) SDN BHD      <- inline, .txt
        Shipper (Principal or Seller) (发货人)             <- block, .docx / .pdf
        APRIL FINE PAPER TRADING ...

    Block form is 36 of 250 attachments and yielded nothing at all until this
    existed. Pass `is_label` (normalize.field_for_label) to enable it: a bare
    line counts as a label only when it names a field we actually want, which
    stops short all-caps values like "SINGAPORE" being read as labels.
    """
    out: list[tuple[str, str, int, int]] = []
    inline_starts = set()
    for m in _LABEL_LINE.finditer(text):
        out.append((m.group(1).strip(), m.group(2).strip(), m.start(2), m.end(2)))
        inline_starts.add(m.start())
    if is_label is None:
        return out

    lines: list[tuple[str, int]] = []
    pos = 0
    for raw in text.split("\n"):
        lines.append((raw, pos))
        pos += len(raw) + 1

    for i, (raw, off) in enumerate(lines):
        if off in inline_starts or not raw.strip():
            continue
        m = _LABEL_ONLY.match(raw)
        if not m or not is_label(m.group(1)):
            continue

        # The value is the next non-empty line -- unless that line is itself a
        # label. Two labels in a row means a table header row, and pairing
        # across it would attach a value from the wrong column.
        j = i + 1
        while j < len(lines) and not lines[j][0].strip():
            j += 1
        if j >= len(lines):
            continue
        nxt, nxt_off = lines[j]
        # "SHIPPER:" with nothing after it is a blank field, not a block label.
        # Pairing it with the next line stole the following field's line whole:
        # shipper became "CONSIGNEE: UAB NOVAKOPA".
        #
        # Only a line naming a *known field* disqualifies it. Rejecting every
        # line with a colon in it was too blunt: an address line reading
        # "NAGAPPA EXPORTSNEW NO : 23, L-BLOCK ..." is a perfectly good value.
        nxt_inline = _LABEL_LINE.match(nxt)
        if nxt_inline and is_label(nxt_inline.group(1)):
            continue
        nxt_label = _LABEL_ONLY.match(nxt)
        if nxt_label and is_label(nxt_label.group(1)):
            continue

        value = nxt.strip()
        start = nxt_off + (len(nxt) - len(nxt.lstrip()))
        out.append((m.group(1).strip(), value, start, start + len(value)))
    return out


def page_images(path: Path) -> list[bytes]:
    """Page images, for a document with no text layer to read."""
    if path.suffix.lower() in IMAGE_SUFFIXES:
        return [path.read_bytes()]
    if path.suffix.lower() != ".pdf":
        return []
    from pypdf import PdfReader

    images: list[bytes] = []
    try:
        for page in PdfReader(str(path)).pages:
            for image in page.images:
                images.append(image.data)
    except Exception:  # noqa: BLE001 - a corrupt file simply has no images to give
        pass
    return images
